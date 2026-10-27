import hashlib
import hmac
import io
import os
import secrets
import sqlite3
from functools import wraps

import openpyxl
from dotenv import load_dotenv
from flask import Flask, jsonify, redirect, render_template, request, send_file, session, url_for
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from openpyxl.styles import Alignment, Font, PatternFill

from app.logging_config import WEB_LOG_PATH, get_logger, setup_logging

ALLOWED_STATUSES = {"new", "in_progress", "done"}

load_dotenv()

logger = get_logger(__name__)

app = Flask(__name__)
SECRET_KEY = os.getenv("SECRET_KEY")
if not SECRET_KEY:
    raise RuntimeError("SECRET_KEY not set in .env")
app.secret_key = SECRET_KEY
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.getenv("SESSION_COOKIE_SECURE", "true").lower()
    in ("1", "true", "yes"),
)

DEMO_MODE = os.environ.get("DEMO_MODE", "").lower() in ("1", "true", "yes")

limiter = Limiter(key_func=get_remote_address, app=app)

FLASK_USERNAME = os.getenv("FLASK_USERNAME")
FLASK_PASSWORD_HASH = os.getenv("FLASK_PASSWORD_HASH")
if not FLASK_USERNAME or not FLASK_PASSWORD_HASH:
    raise RuntimeError("FLASK_USERNAME/FLASK_PASSWORD_HASH not set in .env")
DATABASE_PATH = os.getenv("DATABASE_PATH", "/app/data/database.db")


@app.context_processor
def inject_demo_mode():
    return {"demo_mode": DEMO_MODE}


def get_connection():
    conn = sqlite3.connect(DATABASE_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA busy_timeout=5000")
    return conn


def password_matches(password: str) -> bool:
    digest = hmac.new(SECRET_KEY.encode(), password.encode(), hashlib.sha256).hexdigest()
    return hmac.compare_digest(digest, FLASK_PASSWORD_HASH)


def get_csrf_token() -> str:
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_hex(32)
    return session["csrf_token"]


def csrf_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        token = session.get("csrf_token")
        form_token = request.form.get("csrf_token", "")
        if not token or not hmac.compare_digest(token, form_token):
            return "CSRF token mismatch", 400
        return f(*args, **kwargs)

    return decorated


def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        return f(*args, **kwargs)

    return decorated


@app.route("/login", methods=["GET", "POST"])
@limiter.limit("5 per minute", methods=["POST"])
def login():
    error = None
    if request.method == "POST":
        if request.form["username"] == FLASK_USERNAME and password_matches(
            request.form["password"]
        ):
            session.clear()
            session["logged_in"] = True
            return redirect(url_for("index"))
        error = "Неверный логин или пароль"
    return render_template("login.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/healthz")
def healthz():
    try:
        conn = get_connection()
        conn.execute("SELECT 1")
        conn.close()
    except Exception:
        logger.exception("Healthcheck: БД недоступна")
        return "unavailable", 503
    return "ok", 200


@app.after_request
def set_security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Cache-Control"] = "no-store"
    return response


@app.route("/")
@login_required
def index():
    status_filter = request.args.get("status", "")
    if status_filter not in ALLOWED_STATUSES:
        status_filter = ""
    conn = get_connection()

    if status_filter:
        apps = conn.execute(
            "SELECT * FROM applications WHERE status=? ORDER BY created_at DESC", (status_filter,)
        ).fetchall()
    else:
        apps = conn.execute("SELECT * FROM applications ORDER BY created_at DESC").fetchall()

    stats = {
        "total": conn.execute("SELECT COUNT(*) FROM applications").fetchone()[0],
        "new": conn.execute("SELECT COUNT(*) FROM applications WHERE status='new'").fetchone()[0],
        "in_progress": conn.execute(
            "SELECT COUNT(*) FROM applications WHERE status='in_progress'"
        ).fetchone()[0],
        "done": conn.execute("SELECT COUNT(*) FROM applications WHERE status='done'").fetchone()[0],
    }
    conn.close()

    return render_template(
        "index.html",
        applications=apps,
        stats=stats,
        current_status=status_filter,
        csrf_token=get_csrf_token(),
    )


@app.route("/update_status", methods=["POST"])
@login_required
@csrf_required
def update_status():
    app_id = request.form["id"]
    new_status = request.form["status"]
    if new_status not in ALLOWED_STATUSES:
        return jsonify({"error": f"Invalid status. Allowed: {sorted(ALLOWED_STATUSES)}"}), 400
    status_filter = (
        request.referrer.split("status=")[-1] if "status=" in (request.referrer or "") else ""
    )
    conn = get_connection()
    conn.execute("UPDATE applications SET status=? WHERE id=?", (new_status, app_id))
    conn.commit()
    conn.close()
    return redirect(url_for("index", status=status_filter))


@app.route("/delete", methods=["POST"])
@login_required
@csrf_required
def delete():
    app_id = request.form["id"]
    conn = get_connection()
    conn.execute("DELETE FROM applications WHERE id=?", (app_id,))
    conn.commit()
    conn.close()
    return redirect(request.referrer or url_for("index"))


def sanitize_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _build_workbook(rows) -> io.BytesIO:
    """Собирает Excel-файл с заявками и возвращает его как BytesIO."""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Заявки"

    headers = ["ID", "Имя", "Возраст", "Цель", "Контакт", "Дата", "Статус"]
    header_fill = PatternFill("solid", fgColor="2C3E50")
    header_font = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append(
            [
                row["id"],
                sanitize_cell(row["name"]),
                sanitize_cell(row["age"]),
                sanitize_cell(row["goal"]),
                sanitize_cell(row["contact"]),
                row["created_at"][:16],
                sanitize_cell(row["status"]),
            ]
        )

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 15

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


@app.route("/export")
@login_required
def export():
    status_filter = request.args.get("status", "")
    if status_filter not in ALLOWED_STATUSES:
        status_filter = ""
    conn = get_connection()
    if status_filter:
        rows = conn.execute(
            "SELECT * FROM applications WHERE status=? ORDER BY created_at DESC", (status_filter,)
        ).fetchall()
    else:
        rows = conn.execute("SELECT * FROM applications ORDER BY created_at DESC").fetchall()
    conn.close()

    output = _build_workbook(rows)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="applications.xlsx",
    )


if __name__ == "__main__":
    setup_logging(WEB_LOG_PATH)
    logger.info("Webadmin запущен на порту 5004")
    app.run(host="0.0.0.0", port=5004)  # nosec B104 - намеренно: внешний порт мапится на контейнер
