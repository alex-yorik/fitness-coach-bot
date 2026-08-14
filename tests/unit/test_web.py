import io
import os
import sqlite3

import openpyxl
import pytest

import admin.web as web

TEST_USERNAME = os.environ["FLASK_USERNAME"]
TEST_PASSWORD = "test-admin-password"

CREATE_APPLICATIONS_TABLE = """
CREATE TABLE IF NOT EXISTS applications (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    name TEXT NOT NULL,
    age INTEGER,
    goal TEXT,
    contact TEXT NOT NULL,
    status TEXT NOT NULL DEFAULT 'new'
)
"""


@pytest.fixture
def client(tmp_path, monkeypatch):
    db_path = tmp_path / "web.db"
    monkeypatch.setattr(web, "DATABASE_PATH", str(db_path))
    monkeypatch.setattr(web.limiter, "enabled", False)
    conn = sqlite3.connect(db_path)
    conn.execute(CREATE_APPLICATIONS_TABLE)
    conn.commit()
    conn.close()
    web.app.config["TESTING"] = True
    return web.app.test_client()


def _login(client):
    return client.post("/login", data={"username": TEST_USERNAME, "password": TEST_PASSWORD})


def _insert(client, name="Иван", contact="@ivan"):
    conn = sqlite3.connect(web.DATABASE_PATH)
    conn.execute(
        "INSERT INTO applications (name, age, goal, contact) VALUES (?, ?, ?, ?)",
        (name, 30, "Цель", contact),
    )
    conn.commit()
    row_id = conn.execute("SELECT id FROM applications ORDER BY id DESC LIMIT 1").fetchone()[0]
    conn.close()
    return row_id


def _csrf(client) -> str:
    with client.session_transaction() as sess:
        return sess["csrf_token"]


class TestLogin:
    def test_correct_credentials(self, client):
        resp = _login(client)
        assert resp.status_code == 302
        assert resp.headers["Location"].endswith("/")

        with client.session_transaction() as sess:
            assert sess.get("logged_in") is True

    def test_wrong_password(self, client):
        resp = client.post("/login", data={"username": TEST_USERNAME, "password": "wrong-password"})
        assert resp.status_code == 200
        assert "Неверный логин или пароль" in resp.get_data(as_text=True)

    def test_wrong_username(self, client):
        resp = client.post("/login", data={"username": "hacker", "password": TEST_PASSWORD})
        assert resp.status_code == 200

    def test_get_login_page(self, client):
        assert client.get("/login").status_code == 200

    def test_rate_limit_after_five_attempts(self, client, monkeypatch):
        monkeypatch.setattr(web.limiter, "enabled", True)
        for _ in range(5):
            resp = client.post("/login", data={"username": TEST_USERNAME, "password": "wrong"})
            assert resp.status_code == 200
        resp = client.post("/login", data={"username": TEST_USERNAME, "password": TEST_PASSWORD})
        assert resp.status_code == 429


class TestFailClosed:
    def test_index_redirects_without_session(self, client):
        resp = client.get("/")
        assert resp.status_code == 302
        assert resp.headers["Location"].endswith("/login")

    def test_update_status_redirects_without_session(self, client):
        resp = client.post("/update_status", data={"id": 1, "status": "done"})
        assert resp.status_code == 302
        assert resp.headers["Location"].endswith("/login")

    def test_delete_redirects_without_session(self, client):
        resp = client.post("/delete", data={"id": 1})
        assert resp.status_code == 302
        assert resp.headers["Location"].endswith("/login")

    def test_export_redirects_without_session(self, client):
        resp = client.get("/export")
        assert resp.status_code == 302
        assert resp.headers["Location"].endswith("/login")

    def test_logout_clears_session(self, client):
        _login(client)
        resp = client.get("/logout")
        assert resp.status_code == 302
        with client.session_transaction() as sess:
            assert sess.get("logged_in") is None


class TestHealthz:
    def test_healthz_ok_with_db(self, client):
        resp = client.get("/healthz")
        assert resp.status_code == 200
        assert resp.get_data(as_text=True) == "ok"

    def test_healthz_503_without_db(self, client, monkeypatch):
        def broken_connection():
            raise sqlite3.OperationalError("no such table")

        monkeypatch.setattr(web, "get_connection", broken_connection)
        resp = client.get("/healthz")
        assert resp.status_code == 503
        assert resp.get_data(as_text=True) == "unavailable"

    def test_healthz_not_requires_login(self, client):
        assert client.get("/healthz").status_code == 200


class TestAuthenticated:
    def test_index_with_session(self, client):
        _login(client)
        resp = client.get("/")
        assert resp.status_code == 200
        assert "Заявки" in resp.get_data(as_text=True)

    def test_update_status_with_csrf(self, client):
        _login(client)
        row_id = _insert(client)
        client.get("/")
        csrf = _csrf(client)

        resp = client.post(
            "/update_status",
            data={"id": row_id, "status": "done", "csrf_token": csrf},
        )
        assert resp.status_code == 302

        conn = sqlite3.connect(web.DATABASE_PATH)
        status = conn.execute("SELECT status FROM applications WHERE id=?", (row_id,)).fetchone()[0]
        conn.close()
        assert status == "done"

    def test_delete_with_csrf(self, client):
        _login(client)
        row_id = _insert(client)
        client.get("/")
        csrf = _csrf(client)

        resp = client.post("/delete", data={"id": row_id, "csrf_token": csrf})
        assert resp.status_code == 302

        conn = sqlite3.connect(web.DATABASE_PATH)
        count = conn.execute("SELECT COUNT(*) FROM applications").fetchone()[0]
        conn.close()
        assert count == 0

    def test_csrf_mismatch_rejected(self, client):
        _login(client)
        row_id = _insert(client)

        resp = client.post("/update_status", data={"id": row_id, "status": "done"})
        assert resp.status_code == 400


@pytest.mark.parametrize(
    "value,expected",
    [
        ("=SUM(A1)", "'=SUM(A1)"),
        ("+123", "'+123"),
        ("-123", "'-123"),
        ("@cmd", "'@cmd"),
        ("normal", "normal"),
        ("Иван", "Иван"),
        (None, None),
        (42, 42),
    ],
)
def test_sanitize_cell(value, expected):
    assert web.sanitize_cell(value) == expected


class TestExport:
    def test_export_generates_xlsx(self, client):
        _login(client)
        _insert(client, name="Иван", contact="@ivan")
        _insert(client, name="=2+2", contact="+79990000000")

        resp = client.get("/export")

        assert resp.status_code == 200
        assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert "applications.xlsx" in resp.headers.get("Content-Disposition", "")

        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        ws = wb["Заявки"]
        assert [ws.cell(row=1, column=col).value for col in range(1, 8)] == [
            "ID",
            "Имя",
            "Возраст",
            "Цель",
            "Контакт",
            "Дата",
            "Статус",
        ]
        names = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
        assert "Иван" in names
        assert "'=2+2" in names

    def test_export_with_status_filter(self, client):
        _login(client)
        _insert(client)
        conn = sqlite3.connect(web.DATABASE_PATH)
        conn.execute("UPDATE applications SET status='done'")
        conn.commit()
        conn.close()

        resp = client.get("/export?status=done")
        assert resp.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        ws = wb["Заявки"]
        assert ws.max_row == 2
